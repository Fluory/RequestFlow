"""XLSX parsing with openpyxl (read-only, cached values only).

Segment choice: **one segment per non-empty row**, its non-empty cells joined with `` | `` in
column order (``s{sheet}-r{row}``, 1-based sheet index in workbook order and row number). A row
keeps a line item together (description, quantity, unit), which the model needs to read it as one
position; the quote a field cites is still a substring of that row. The locator names the sheet,
the row and the range of the non-empty cells (``A7:D7``; one cell: ``B7``).

Why not docling: its XLSX backend emits tables and text without cell provenance, and this service
needs stable row/cell locators. docling depends on openpyxl already.

Safety: the zip is checked first (``ooxml.open_package``); openpyxl parses XML through defusedxml
(installed), reads formulas' *cached* values only (``data_only=True``: formula text is never sent,
nothing is computed), ignores external links and never touches macros. The worksheet's declared
dimension is discarded (``reset_dimensions``) so a forged ``A1:XFD1048576`` cannot make openpyxl
pad millions of empty cells. Caps: ``MAX_SHEETS``, ``MAX_ROWS`` non-empty rows, ``MAX_CELLS``
visited cells -> ``DocumentTooLongError``.
"""

from __future__ import annotations

import datetime as dt
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from requestflow_ai.parsing.budget import ParseBudget
from requestflow_ai.parsing.errors import DocumentParseError, DocumentTooLongError
from requestflow_ai.parsing.ooxml import open_package
from requestflow_ai.parsing.segments import Segment, XlsxLocator

MAX_SHEETS = 50
MAX_ROWS = 10_000
MAX_CELLS = 500_000


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return " ".join(str(value).split())


def parse_xlsx(data: bytes, budget: ParseBudget | None = None) -> list[Segment]:
    open_package(data, budget).close()
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise DocumentParseError("could not open workbook") from exc
    try:
        return _segments(workbook)
    except (DocumentParseError, DocumentTooLongError):
        raise
    except Exception as exc:
        raise DocumentParseError("could not read workbook") from exc
    finally:
        workbook.close()


def _segments(workbook: Any) -> list[Segment]:
    sheets = workbook.worksheets
    if len(sheets) > MAX_SHEETS:
        raise DocumentTooLongError(f"workbook has more than {MAX_SHEETS} sheets")
    segments: list[Segment] = []
    visited = 0
    for sheet_no, sheet in enumerate(sheets, start=1):
        if not hasattr(sheet, "iter_rows"):
            continue  # chart sheets have no cells
        sheet.reset_dimensions()
        for row in sheet.iter_rows():
            visited += len(row)
            if visited > MAX_CELLS:
                raise DocumentTooLongError(f"workbook has more than {MAX_CELLS} cells")
            cells = [
                (cell.row, cell.column, text)
                for cell in row
                if getattr(cell, "row", None) is not None and (text := _cell_text(cell.value))
            ]
            if not cells:
                continue
            if len(segments) >= MAX_ROWS:
                raise DocumentTooLongError(f"workbook has more than {MAX_ROWS} rows with text")
            row_no = int(cells[0][0])
            first = f"{get_column_letter(cells[0][1])}{row_no}"
            last = f"{get_column_letter(cells[-1][1])}{row_no}"
            segments.append(
                Segment(
                    id=f"s{sheet_no}-r{row_no}",
                    text=" | ".join(text for _, _, text in cells),
                    locator=XlsxLocator(
                        sheet=str(sheet.title),
                        row=row_no,
                        cell_range=first if first == last else f"{first}:{last}",
                    ),
                )
            )
    return segments
