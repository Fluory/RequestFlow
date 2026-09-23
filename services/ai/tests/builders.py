"""Build small synthetic XLSX, DOCX and Outlook .msg documents in-test (all content invented).

XLSX and DOCX are written with openpyxl and python-docx. Nothing in the dependency set can *write*
an Outlook .msg (olefile and python-oxmsg only read), so ``build_msg`` contains a minimal Compound
File Binary (CFB v3) writer: 512-byte sectors, one FAT (no DIFAT sectors), a mini stream for
streams below 4096 bytes, and the MSG property streams python-oxmsg reads ([MS-OXMSG] 2.4).
It is test code only; the service never writes documents.
"""

from __future__ import annotations

import struct
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from openpyxl import Workbook

# --- XLSX / DOCX --------------------------------------------------------------------------------

CellValue = str | int | float | None


def build_xlsx(sheets: Mapping[str, Sequence[Sequence[CellValue]]]) -> bytes:
    """One sheet per key; rows start at A1. ``None`` leaves a cell empty."""
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@dataclass
class DocxTable:
    rows: Sequence[Sequence[str]]
    merge_first_row: bool = False  # merge the first two cells of row 1 (a horizontal span)


def build_docx(blocks: Sequence[str | DocxTable]) -> bytes:
    """Paragraphs (``str``, may be empty) and tables in body order."""
    document = Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
            continue
        width = max(len(row) for row in block.rows)
        table = document.add_table(rows=len(block.rows), cols=width)
        for r, row in enumerate(block.rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
        if block.merge_first_row:
            merged = table.cell(0, 0).merge(table.cell(0, 1))
            merged.text = block.rows[0][0]
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- CFB writer ---------------------------------------------------------------------------------

_SECTOR = 512
_MINI = 64
_CUTOFF = 4096
_FREE = 0xFFFFFFFF
_END = 0xFFFFFFFE
_FATSECT = 0xFFFFFFFD
_NOSTREAM = 0xFFFFFFFF

Tree = Mapping[str, "bytes | Tree"]


@dataclass
class _Entry:
    name: str
    kind: int  # 1 storage, 2 stream, 5 root
    data: bytes = b""
    children: list[int] = field(default_factory=list[int])
    left: int = _NOSTREAM
    right: int = _NOSTREAM
    child: int = _NOSTREAM
    start: int = _END
    size: int = 0


def _flatten(tree: Tree) -> list[_Entry]:
    entries = [_Entry("Root Entry", 5)]

    def add(parent: int, node: Tree) -> None:
        for name, value in node.items():
            index = len(entries)
            if isinstance(value, bytes):
                entries.append(_Entry(name, 2, data=value, size=len(value)))
            else:
                entries.append(_Entry(name, 1))
                add(index, value)
            entries[parent].children.append(index)

    add(0, tree)
    return entries


def _link_children(entries: list[_Entry]) -> None:
    """Children of a storage as a balanced binary tree ordered by (length, upper-case name)."""

    def build(ids: list[int]) -> int:
        if not ids:
            return _NOSTREAM
        middle = len(ids) // 2
        node = entries[ids[middle]]
        node.left = build(ids[:middle])
        node.right = build(ids[middle + 1 :])
        return ids[middle]

    for entry in entries:
        ordered = sorted(entry.children, key=lambda i: (len(entries[i].name), entries[i].name.upper()))
        entry.child = build(ordered)


def _chain(fat: list[int], start: int, count: int) -> None:
    for offset in range(count):
        fat[start + offset] = start + offset + 1 if offset < count - 1 else _END


def _sectors(size: int, unit: int) -> int:
    return (size + unit - 1) // unit


def build_cfb(tree: Tree) -> bytes:
    entries = _flatten(tree)
    _link_children(entries)

    # Mini stream: every stream below the cutoff, in 64-byte mini sectors.
    mini_fat: list[int] = []
    mini_stream = bytearray()
    for entry in entries:
        if entry.kind == 2 and 0 < entry.size < _CUTOFF:
            count = _sectors(entry.size, _MINI)
            entry.start = len(mini_fat)
            mini_fat.extend([0] * count)
            _chain(mini_fat, entry.start, count)
            mini_stream += entry.data.ljust(count * _MINI, b"\x00")
    large = [e for e in entries if e.kind == 2 and e.size >= _CUTOFF]

    dir_sectors = _sectors(len(entries) * 128, _SECTOR)
    minifat_sectors = _sectors(len(mini_fat) * 4, _SECTOR)
    ministream_sectors = _sectors(len(mini_stream), _SECTOR)
    large_sectors = sum(_sectors(e.size, _SECTOR) for e in large)
    payload = dir_sectors + minifat_sectors + ministream_sectors + large_sectors
    fat_sectors = 1
    while fat_sectors * (_SECTOR // 4) < fat_sectors + payload:
        fat_sectors += 1
    assert fat_sectors <= 109, "fixture too large for a header-only DIFAT"

    fat = [_FREE] * (fat_sectors * (_SECTOR // 4))
    for index in range(fat_sectors):
        fat[index] = _FATSECT
    cursor = fat_sectors
    dir_start = cursor
    _chain(fat, cursor, dir_sectors)
    cursor += dir_sectors
    minifat_start = cursor if minifat_sectors else _END
    _chain(fat, cursor, minifat_sectors)
    cursor += minifat_sectors
    root = entries[0]
    root.start = cursor if ministream_sectors else _END
    root.size = len(mini_stream)
    _chain(fat, cursor, ministream_sectors)
    cursor += ministream_sectors
    body = bytearray()
    for entry in large:
        count = _sectors(entry.size, _SECTOR)
        entry.start = cursor
        _chain(fat, cursor, count)
        cursor += count
        body += entry.data.ljust(count * _SECTOR, b"\x00")

    directory = bytearray()
    for entry in entries:
        name = entry.name.encode("utf-16-le")
        assert len(name) <= 62, "CFB names are at most 31 characters"
        directory += struct.pack(
            "<64sHBBIII16sIQQIQ",
            name,
            len(name) + 2,
            entry.kind,
            1,  # black
            entry.left,
            entry.right,
            entry.child,
            b"\x00" * 16,
            0,
            0,
            0,
            entry.start,
            entry.size,
        )
    empty = struct.pack("<64sHBBIII16sIQQIQ", b"", 0, 0, 0, _NOSTREAM, _NOSTREAM, _NOSTREAM,
                        b"\x00" * 16, 0, 0, 0, 0, 0)  # fmt: skip
    while len(directory) % _SECTOR:
        directory += empty

    difat = list(range(fat_sectors)) + [_FREE] * (109 - fat_sectors)
    header = struct.pack(
        "<8s16sHHHHH6sIIIIIIIII",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
        b"\x00" * 16,
        0x003E,
        3,
        0xFFFE,
        9,
        6,
        b"\x00" * 6,
        0,
        fat_sectors,
        dir_start,
        0,
        _CUTOFF,
        minifat_start,
        minifat_sectors,
        _END,
        0,
    ) + struct.pack("<109I", *difat)

    out = bytearray(header)
    out += struct.pack(f"<{len(fat)}I", *fat)
    out += directory
    minifat_bytes = struct.pack(f"<{len(mini_fat)}I", *mini_fat) if mini_fat else b""
    out += minifat_bytes.ljust(minifat_sectors * _SECTOR, b"\xff")
    out += bytes(mini_stream).ljust(ministream_sectors * _SECTOR, b"\x00")
    out += body
    return bytes(out)


# --- MSG ----------------------------------------------------------------------------------------

_PT_LONG = 0x0003
_PT_OBJECT = 0x000D
_PT_UNICODE = 0x001F
_PT_BINARY = 0x0102
_ATTACH_BY_VALUE = 1
_ATTACH_EMBEDDED_MSG = 5


@dataclass
class MsgAttachment:
    name: str
    data: bytes | None = None  # by value
    mime_type: str | None = None
    embedded: MsgSpec | None = None  # an attached Outlook item (stored as a sub-storage)


@dataclass
class MsgSpec:
    subject: str = ""
    sender_name: str = ""
    sender_email: str = ""
    body: str | None = None
    html_body: str | None = None
    attachments: list[MsgAttachment] = field(default_factory=list[MsgAttachment])


def _properties(header: bytes, props: Sequence[tuple[int, int, bytes | int]]) -> Tree:
    streams: dict[str, bytes | Tree] = {}
    table = bytearray(header)
    for pid, ptyp, value in props:
        tag = (pid << 16) | ptyp
        if isinstance(value, int):
            table += struct.pack("<IIIxxxx", tag, 6, value)
            continue
        size = len(value) + (2 if ptyp == _PT_UNICODE else 0)
        table += struct.pack("<IIIxxxx", tag, 6, size)
        streams[f"__substg1.0_{pid:04X}{ptyp:04X}"] = value
    streams["__properties_version1.0"] = bytes(table)
    return streams


def _unicode(value: str) -> bytes:
    return value.encode("utf-16-le")


def _message_tree(spec: MsgSpec, embedded: bool) -> Tree:
    props: list[tuple[int, int, bytes | int]] = [(0x001A, _PT_UNICODE, _unicode("IPM.Note"))]
    if spec.subject:
        props.append((0x0037, _PT_UNICODE, _unicode(spec.subject)))
    if spec.sender_name:
        props.append((0x0C1A, _PT_UNICODE, _unicode(spec.sender_name)))
    if spec.sender_email:
        props.append((0x0C1F, _PT_UNICODE, _unicode(spec.sender_email)))
    if spec.body is not None:
        props.append((0x1000, _PT_UNICODE, _unicode(spec.body)))
    if spec.html_body is not None:
        props.append((0x3FDE, _PT_LONG, 65001))  # internet code page: UTF-8
        props.append((0x1013, _PT_BINARY, spec.html_body.encode("utf-8")))
    count = len(spec.attachments)
    # Root header: 8 reserved, next recipient id, next attachment id, recipient/attachment count,
    # then 8 reserved (32 bytes); an embedded message has no trailing reserved bytes (24 bytes).
    header = struct.pack("<8xIIII", 0, count, 0, count) + (b"" if embedded else b"\x00" * 8)
    tree: dict[str, bytes | Tree] = dict(_properties(header, props))
    for index, attachment in enumerate(spec.attachments):
        attach_props: list[tuple[int, int, bytes | int]] = [
            (0x3707, _PT_UNICODE, _unicode(attachment.name)),
        ]
        storage: dict[str, bytes | Tree]
        if attachment.embedded is not None:
            attach_props.append((0x3705, _PT_LONG, _ATTACH_EMBEDDED_MSG))
            attach_props.append((0x3701, _PT_OBJECT, 0))
            storage = dict(_properties(b"\x00" * 8, attach_props))
            storage["__substg1.0_3701000D"] = _message_tree(attachment.embedded, embedded=True)
        else:
            attach_props.append((0x3705, _PT_LONG, _ATTACH_BY_VALUE))
            if attachment.mime_type:
                attach_props.append((0x370E, _PT_UNICODE, _unicode(attachment.mime_type)))
            if attachment.data is not None:
                attach_props.append((0x3701, _PT_BINARY, attachment.data))
            storage = dict(_properties(b"\x00" * 8, attach_props))
        tree[f"__attach_version1.0_#{index:08X}"] = storage
    return tree


def build_msg(spec: MsgSpec) -> bytes:
    return build_cfb(_message_tree(spec, embedded=False))
